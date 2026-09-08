"""Read the retained planning workbooks; never modify them.

Run with the bundled Python runtime. The JSON is provenance, not a second
editable plan; maintain the public syllabus-data.js once imported.
"""
import json
import hashlib
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[3]
planning = Path(__file__).parent
primary = openpyxl.load_workbook(planning / 'Syllabus planner.xlsx', data_only=True)
coverage = openpyxl.load_workbook(planning / 'A-Level_Macro_International_32-Lesson_Coverage_Planner.xlsx', data_only=True)
estimates = {}
for row in coverage['Detailed Checklist'].iter_rows(min_row=8, values_only=True):
    if row[4]:
        estimates[row[4]] = row[9]

sections = []
for ws in primary.worksheets:
    section = {'id': ws.title.split()[0], 'title': ws['A1'].value.replace(' (A Level)', ''), 'sheet': ws.title, 'points': []}
    topic_title = ''
    for row in ws.iter_rows(min_row=3):
        values = [c.value for c in row]
        if not values[2]:
            continue
        if values[1]:
            topic_title = values[1]
        section['points'].append({
            'code': values[2], 'topic': values[0], 'topicTitle': topic_title,
            'wording': values[3], 'bullets': values[4] or '',
            'sourceAllocation': values[5],
            'allocation': values[5] if values[5] is not None else estimates[values[2]],
            'provisional': values[5] is None,
            'sourceRange': f'C{row[0].row}:F{row[0].row}',
        })
    sections.append(section)

output = planning / 'syllabus-workbook-source.json'
sources = [{'name': name, 'sha256': hashlib.sha256((planning / name).read_bytes()).hexdigest()} for name in ['Syllabus planner.xlsx', 'A-Level_Macro_International_32-Lesson_Coverage_Planner.xlsx']]
output.write_text(json.dumps({'sources': sources, 'sections': sections}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
for section in sections:
    print(f"Section {section['id']}: {len(section['points'])} statements, {sum(p['allocation'] for p in section['points'])} lessons")
